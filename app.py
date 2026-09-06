"""
Streamlit Web Application for E-Commerce Product Data Web Scraper.
"""

from io import BytesIO
import streamlit as st
import pandas as pd
import plotly.express as px
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from scraper.scraper_manager import ScraperManager
from scraper.data_processor import clean_data
from scraper.analyzer import (
    average_price,
    average_rating,
    available_products,
    unavailable_products,
    most_common_category,
    highest_rated_products,
    lowest_priced_products,
    most_expensive_products
)
from scraper.base_scraper import ApiNotConfiguredError, ApiError
from scraper.logger import logger

# ============================================================
# PAGE CONFIGURATION & STYLING
# ============================================================
st.set_page_config(
    page_title="E-Commerce Product Web Scraper",
    page_icon="🛒",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Styling for polished appearance
st.markdown(
    """
    <style>
    .main-header {
        font-size: 2.2rem;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 0.2rem;
    }
    .sub-header {
        font-size: 1.05rem;
        color: #4B5563;
        margin-bottom: 1.5rem;
    }
    .metric-card {
        background-color: #F8FAFC;
        border-radius: 8px;
        padding: 12px;
        border: 1px solid #E2E8F0;
    }
    .stDownloadButton button {
        width: 100%;
        border-radius: 6px;
        font-weight: 600;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# HEADER
# ============================================================
st.markdown('<div class="main-header">🛒 E-Commerce Product Data Web Scraper</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="sub-header">'
    'Scrape, clean, standardize, and analyze structured e-commerce product data with automated compliance controls and multi-format exports.'
    '</div>',
    unsafe_allow_html=True
)

# ============================================================
# INITIALIZE MANAGERS & SESSION STATE
# ============================================================
manager = ScraperManager()

if "products" not in st.session_state:
    st.session_state.products = []

if "search_info" not in st.session_state:
    st.session_state.search_info = None

# ============================================================
# SIDEBAR CONTROLS
# ============================================================
st.sidebar.header("⚙️ Scraper Controls")

website_option = st.sidebar.selectbox(
    "Target Source",
    options=["Test Site", "Amazon", "Flipkart", "Alibaba"],
    index=0,
    help="Select 'Test Site' for a guaranteed safe live scraping demo without API credentials."
)

query_input = st.sidebar.text_input(
    "Search Query",
    value="light",
    placeholder="e.g. light, art, music, fiction...",
    help="Enter keyword to search product catalog."
)

max_products_input = st.sidebar.slider(
    "Maximum Products",
    min_value=1,
    max_value=100,
    value=10,
    step=1,
    help="Maximum number of products to retrieve."
)

search_clicked = st.sidebar.button("🔎 Search Products", use_container_width=True)

st.sidebar.markdown("---")
st.sidebar.markdown(
    """
    **ℹ️ Data Access Mode**
    - **Test Site**: Public demonstration site (Books to Scrape) with full HTTP scraping, robots.txt compliance, and pagination.
    - **Amazon / Flipkart / Alibaba**: Configured for authorized APIs. Requires credentials in `.env`.
    """
)

# ============================================================
# EXECUTE SCRAPING UPON BUTTON CLICK
# ============================================================
if search_clicked:
    clean_query = query_input.strip()

    if not clean_query:
        st.sidebar.error("Please enter a valid search keyword.")
    else:
        site_key = manager.normalize_name(website_option)
        scraper = manager.get_scraper(site_key)

        # Check if requested API requires configuration
        if site_key != "test" and not scraper.is_configured():
            st.warning(
                f"⚠️ **{website_option} API is not configured.**\n\n"
                f"Please configure the required authorized API credentials via `{scraper.env_var}` in your `.env` file.\n\n"
                f"💡 *Tip: Switch to **Test Site** in the sidebar to test the complete live scraping and data processing pipeline without credentials.*"
            )
            st.session_state.products = []
            st.session_state.search_info = None
        else:
            with st.spinner(f"Querying {website_option} for '{clean_query}' (up to {max_products_input} items)..."):
                try:
                    logger.info(f"UI initiated search: website='{website_option}', query='{clean_query}', max={max_products_input}")
                    raw_products = scraper.search_products(clean_query, max_products=max_products_input)

                    if not raw_products:
                        st.info(f"No products found matching '{clean_query}' on {website_option}.")
                        st.session_state.products = []
                        st.session_state.search_info = None
                    else:
                        raw_df = pd.DataFrame(raw_products)
                        cleaned_df = clean_data(raw_df)

                        if cleaned_df.empty:
                            st.warning("Products were found, but none passed data validation standards.")
                            st.session_state.products = []
                            st.session_state.search_info = None
                        else:
                            st.session_state.products = cleaned_df.to_dict(orient="records")
                            st.session_state.search_info = {
                                "website": website_option,
                                "query": clean_query,
                                "count": len(cleaned_df)
                            }
                            st.success(f"Successfully collected and validated {len(cleaned_df)} products from {website_option}!")

                except ApiNotConfiguredError as e:
                    st.warning(f"⚠️ {e}")
                    st.session_state.products = []
                    st.session_state.search_info = None
                except ApiError as e:
                    st.error(f"{e}")
                    st.session_state.products = []
                    st.session_state.search_info = None
                except Exception as e:
                    logger.exception(f"UI scraping error: {e}")
                    st.error(f"Scraping operation encountered an issue: {e}")
                    st.session_state.products = []
                    st.session_state.search_info = None

# ============================================================
# RENDER RESULTS
# ============================================================
if st.session_state.products:
    df = pd.DataFrame(st.session_state.products)

    # --------------------------------------------------------
    # 1. SUMMARY METRICS
    # --------------------------------------------------------
    st.markdown("### 📊 Summary Metrics")
    col1, col2, col3, col4 = st.columns(4)

    total_count = len(df)
    avg_p = average_price(df)
    avg_r = average_rating(df)
    in_stock_cnt = available_products(df)

    with col1:
        st.metric(label="Products Found", value=total_count)
    with col2:
        st.metric(label="Average Price", value=f"£{avg_p:.2f}" if avg_p > 0 else "N/A")
    with col3:
        st.metric(label="Average Rating", value=f"{avg_r:.1f} ⭐" if avg_r > 0 else "N/A")
    with col4:
        st.metric(label="In Stock", value=f"{in_stock_cnt} / {total_count}")

    st.markdown("---")

    # --------------------------------------------------------
    # 2. PRODUCT TABLE
    # --------------------------------------------------------
    st.markdown("### 📦 Cleaned Product Catalog")

    # Format display columns
    display_df = df.copy()
    columns_to_show = ["name", "price", "rating", "availability", "category", "source", "url", "description"]
    display_df = display_df[[c for c in columns_to_show if c in display_df.columns]]

    # Configure clickable link and friendly labels
    st.dataframe(
        display_df,
        column_config={
            "name": st.column_config.TextColumn("Product Name", width="medium"),
            "price": st.column_config.NumberColumn("Price (£)", format="£%.2f"),
            "rating": st.column_config.NumberColumn("Rating (⭐)", format="%.1f"),
            "availability": st.column_config.TextColumn("Availability"),
            "category": st.column_config.TextColumn("Category"),
            "source": st.column_config.TextColumn("Source"),
            "url": st.column_config.LinkColumn("Product Link", display_text="Visit Page"),
            "description": st.column_config.TextColumn("Description", width="large")
        },
        use_container_width=True,
        hide_index=True
    )

    st.markdown("---")

    # --------------------------------------------------------
    # 3. VISUALIZATIONS & ANALYSIS
    # --------------------------------------------------------
    st.markdown("### 📈 Visualizations & Data Analysis")

    viz_col1, viz_col2 = st.columns(2)

    with viz_col1:
        # Price Distribution
        if "price" in df.columns and df["price"].dropna().count() > 0:
            fig_price = px.histogram(
                df.dropna(subset=["price"]),
                x="price",
                nbins=12,
                title="💰 Price Distribution",
                labels={"price": "Price (£)"},
                color_discrete_sequence=["#3B82F6"]
            )
            fig_price.update_layout(bargap=0.1, template="plotly_white")
            st.plotly_chart(fig_price, use_container_width=True)

        # Availability Pie Chart
        if "availability" in df.columns and not df["availability"].dropna().empty:
            avail_counts = df["availability"].value_counts().reset_index()
            avail_counts.columns = ["Availability", "Count"]
            fig_avail = px.pie(
                avail_counts,
                names="Availability",
                values="Count",
                title="📦 Availability Ratio",
                color_discrete_sequence=["#10B981", "#EF4444", "#6B7280"],
                hole=0.4
            )
            fig_avail.update_layout(template="plotly_white")
            st.plotly_chart(fig_avail, use_container_width=True)

    with viz_col2:
        # Rating Distribution
        if "rating" in df.columns and df["rating"].dropna().count() > 0:
            rating_counts = df["rating"].dropna().value_counts().sort_index().reset_index()
            rating_counts.columns = ["Rating", "Count"]
            fig_rating = px.bar(
                rating_counts,
                x="Rating",
                y="Count",
                title="⭐ Rating Breakdown",
                labels={"Rating": "Stars", "Count": "Number of Products"},
                color_discrete_sequence=["#F59E0B"]
            )
            fig_rating.update_layout(xaxis=dict(dtick=1), template="plotly_white")
            st.plotly_chart(fig_rating, use_container_width=True)

        # Category Distribution
        if "category" in df.columns:
            valid_cats = df[df["category"].notna() & (df["category"] != "Unknown")]
            if not valid_cats.empty:
                cat_counts = valid_cats["category"].value_counts().head(8).reset_index()
                cat_counts.columns = ["Category", "Count"]
                fig_cat = px.bar(
                    cat_counts,
                    x="Count",
                    y="Category",
                    orientation="h",
                    title="🏷️ Top Categories",
                    labels={"Category": "Category", "Count": "Count"},
                    color_discrete_sequence=["#8B5CF6"]
                )
                fig_cat.update_layout(yaxis=dict(autorange="reversed"), template="plotly_white")
                st.plotly_chart(fig_cat, use_container_width=True)

    st.markdown("---")

    # --------------------------------------------------------
    # 4. KEY INSIGHTS
    # --------------------------------------------------------
    st.markdown("### 💡 Key Insights")
    ins_col1, ins_col2, ins_col3 = st.columns(3)

    most_exp = most_expensive_products(df, top_n=1)
    cheapest = lowest_priced_products(df, top_n=1)
    highest_r = highest_rated_products(df, top_n=1)
    common_cat = most_common_category(df)
    unavail_cnt = unavailable_products(df)

    with ins_col1:
        st.markdown("**💰 Pricing Highlights**")
        if not most_exp.empty:
            st.info(f"**Most Expensive:**\n\n{most_exp.iloc[0]['name']}\n\n**Price:** £{most_exp.iloc[0]['price']:.2f}")
        if not cheapest.empty:
            st.success(f"**Cheapest:**\n\n{cheapest.iloc[0]['name']}\n\n**Price:** £{cheapest.iloc[0]['price']:.2f}")

    with ins_col2:
        st.markdown("**⭐ Quality & Ratings**")
        if not highest_r.empty:
            st.info(f"**Highest Rated:**\n\n{highest_r.iloc[0]['name']}\n\n**Rating:** {highest_r.iloc[0]['rating']} ⭐")
        st.write(f"**Average Rating:** {avg_r:.2f} ⭐")

    with ins_col3:
        st.markdown("**📦 Inventory & Categories**")
        st.write(f"**Dominant Category:** {common_cat}")
        st.write(f"**In-Stock Items:** {in_stock_cnt}")
        st.write(f"**Out-of-Stock Items:** {unavail_cnt}")

    st.markdown("---")

    # --------------------------------------------------------
    # 5. DATA EXPORT / DOWNLOAD BUTTONS
    # --------------------------------------------------------
    st.markdown("### ⬇️ Download Cleaned Dataset")
    down_col1, down_col2 = st.columns(2)

    with down_col1:
        csv_bytes = df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
        st.download_button(
            label="📄 Download CSV",
            data=csv_bytes,
            file_name="products.csv",
            mime="text/csv",
            use_container_width=True
        )

    with down_col2:
        excel_buffer = BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Products")
            ws = writer.sheets["Products"]

            # Style header row
            h_font = Font(name="Calibri", size=11, bold=True, color="1F2937")
            h_fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
            for col_idx in range(1, len(df.columns) + 1):
                c = ws.cell(row=1, column=col_idx)
                c.font = h_font
                c.fill = h_fill
                c.alignment = Alignment(horizontal="center", vertical="center")

            # Column widths
            for idx, col in enumerate(df.columns, start=1):
                letter = get_column_letter(idx)
                val_len = df[col].dropna().astype(str).str.len().max() if not df.empty else 10
                ws.column_dimensions[letter].width = min(max(int(val_len or 10) + 3, 12), 45)

        excel_bytes = excel_buffer.getvalue()
        st.download_button(
            label="📊 Download Excel",
            data=excel_bytes,
            file_name="products.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

else:
    # Empty State Guidance
    st.info(
        "👋 **Welcome! Get started by searching:**\n\n"
        "1. Select **Test Site** in the sidebar (or configure your authorized API keys for Amazon/Flipkart/Alibaba).\n"
        "2. Enter a keyword (e.g. `light`, `poetry`, `art`) and choose the maximum products.\n"
        "3. Click **🔎 Search Products** to scrape, clean, visualize, and export the dataset."
    )