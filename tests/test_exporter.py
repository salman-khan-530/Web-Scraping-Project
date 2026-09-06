"""
Unit tests for data exporter (CSV and Excel export).
"""

import unittest
import tempfile
from pathlib import Path
import pandas as pd
import openpyxl

from scraper.exporter import export_to_csv, export_to_excel, export_data


class TestExporter(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)
        self.sample_df = pd.DataFrame([
            {
                "name": "Test Book",
                "price": 29.99,
                "rating": 4.0,
                "availability": "In Stock",
                "url": "https://example.com/test-book",
                "category": "Fiction",
                "description": "A great story.",
                "source": "Test Site"
            }
        ])

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_export_to_csv(self):
        csv_file = self.dir_path / "subfolder" / "products.csv"
        export_to_csv(self.sample_df, csv_file)

        self.assertTrue(csv_file.exists())
        read_df = pd.read_csv(csv_file)
        self.assertEqual(len(read_df), 1)
        self.assertEqual(read_df.iloc[0]["name"], "Test Book")
        self.assertEqual(read_df.iloc[0]["price"], 29.99)
        self.assertIn("category", read_df.columns)
        self.assertIn("source", read_df.columns)

    def test_export_to_excel(self):
        excel_file = self.dir_path / "subfolder" / "products.xlsx"
        export_to_excel(self.sample_df, excel_file)

        self.assertTrue(excel_file.exists())
        wb = openpyxl.load_workbook(excel_file)
        self.assertIn("Products", wb.sheetnames)
        ws = wb["Products"]
        # Check header styling
        self.assertTrue(ws.cell(row=1, column=1).font.bold)
        # Check data
        self.assertEqual(ws.cell(row=2, column=1).value, "Test Book")

    def test_export_data(self):
        csv_file = self.dir_path / "both.csv"
        excel_file = self.dir_path / "both.xlsx"
        export_data(self.sample_df, csv_file, excel_file)

        self.assertTrue(csv_file.exists())
        self.assertTrue(excel_file.exists())


if __name__ == "__main__":
    unittest.main()
