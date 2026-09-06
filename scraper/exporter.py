"""
Export functions for CSV and Excel formats with formatting.
"""

from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from scraper.logger import logger

COLUMNS_ORDER = [
    "name",
    "price",
    "rating",
    "availability",
    "url",
    "category",
    "description",
    "source"
]


def _prepare_df(df):
    """Ensure standard column ordering."""
    existing_cols = [c for c in COLUMNS_ORDER if c in df.columns]
    other_cols = [c for c in df.columns if c not in COLUMNS_ORDER]
    return df[existing_cols + other_cols].copy()


def export_to_csv(df, file_path):
    """
    Export DataFrame to CSV with UTF-8 encoding.

    Args:
        df (pandas.DataFrame): Cleaned product data.
        file_path (str or Path): Target CSV path.
    """
    if df is None or df.empty:
        logger.warning("Attempted to export empty DataFrame to CSV.")

    target_path = Path(file_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    ordered_df = _prepare_df(df) if df is not None and not df.empty else pd.DataFrame(columns=COLUMNS_ORDER)
    ordered_df.to_csv(target_path, index=False, encoding="utf-8-sig")
    logger.info(f"CSV file saved successfully ({len(ordered_df)} records): {target_path}")


def export_to_excel(df, file_path):
    """
    Export DataFrame to Excel with styled header and auto-sized column widths.

    Args:
        df (pandas.DataFrame): Cleaned product data.
        file_path (str or Path): Target Excel path.
    """
    target_path = Path(file_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    ordered_df = _prepare_df(df) if df is not None and not df.empty else pd.DataFrame(columns=COLUMNS_ORDER)

    with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
        sheet_name = "Products"
        ordered_df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        # Style headers: Bold, Dark text on subtle gray/blue header, centered
        header_font = Font(name="Calibri", size=11, bold=True, color="1F2937")
        header_fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

        for col_num in range(1, len(ordered_df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-adjust column widths
        for col_idx, col_name in enumerate(ordered_df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col_name))
            if not ordered_df.empty:
                col_values = ordered_df[col_name].dropna().astype(str)
                if not col_values.empty:
                    val_max = col_values.str.len().max()
                    max_len = max(max_len, int(val_max))
            # Bound width between 12 and 45 characters
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    logger.info(f"Excel file saved successfully ({len(ordered_df)} records): {target_path}")


def export_data(df, csv_path, excel_path):
    """Export to both CSV and Excel."""
    export_to_csv(df, csv_path)
    export_to_excel(df, excel_path)